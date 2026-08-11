#!/usr/bin/env python3
import json
import sys
from pathlib import Path


HEADERS = ("姓名", "性别", "证件号码", "手机号码", "邮箱", "科目", "准考证号", "考试状态")


def text(value):
    if value is None:
        return ""
    return str(value)


def export_not_started_candidates(payload_path, output_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:
        raise RuntimeError("当前 Python 环境缺少 openpyxl，请先安装依赖") from exc

    payload = json.loads(Path(payload_path).read_text("utf-8"))
    session = payload.get("session") or {}
    rows = payload.get("rows") or []
    if not text(session.get("session_id")).strip():
        return {"ok": False, "errors": ["缺少 session_id"]}
    if not rows:
        return {"ok": False, "errors": ["当前场次没有未开考考生"]}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "未开考考生"
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(
            [
                text(row.get("name")),
                text(row.get("gender")),
                text(row.get("identity_id")),
                text(row.get("mobile")),
                text(row.get("email")),
                text(row.get("course")),
                text(row.get("permit")),
                text(row.get("exam_status")),
            ]
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 16,
        "B": 10,
        "C": 24,
        "D": 16,
        "E": 28,
        "F": 36,
        "G": 18,
        "H": 14,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(output_path)
    return {"ok": True, "path": str(output_path), "rows": len(rows), "errors": []}


def main():
    if len(sys.argv) != 3:
        print(json.dumps({"ok": False, "errors": ["用法：not_started_candidate_exporter.py payload.json output.xlsx"]}, ensure_ascii=False))
        return 1
    try:
        result = export_not_started_candidates(sys.argv[1], sys.argv[2])
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result.get("ok") else 2
    except Exception as exc:
        print(json.dumps({"ok": False, "errors": [str(exc)]}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
